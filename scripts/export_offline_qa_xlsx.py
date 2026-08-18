#!/usr/bin/env python3
"""Build Excel workbook from offline Q&A catalog JSON (V4)."""

import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
JSON_IN = ROOT / "docs/offline-qa-catalog-v4.json"
XLSX_OUT = ROOT / "docs/V4_Offline_QA_Catalog.xlsx"

HEADERS = [
    "ID",
    "Topic",
    "Trigger keywords",
    "Sample question",
    "User ID",
    "Profile (simulated)",
    "Assignment type",
    "Host country",
    "Full response (complete)",
    "Follow-up 1",
    "Follow-up 2",
    "Follow-up 3",
    "Validation status",
    "Notes / edits",
]

COL_WIDTHS = [6, 22, 28, 42, 14, 36, 12, 16, 80, 36, 36, 36, 14, 36]


def style_header(ws):
    fill = PatternFill("solid", fgColor="0F3D4C")
    font = Font(bold=True, color="FFFFFF", size=11)
    border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )
    for col, title in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[1].height = 28


def main():
    if not JSON_IN.exists():
        raise SystemExit(f"Missing {JSON_IN} — run: node scripts/export_offline_qa_catalog.mjs")

    data = json.loads(JSON_IN.read_text(encoding="utf-8"))
    rows = data.get("rows", [])

    wb = Workbook()
    ws = wb.active
    ws.title = "Offline Q&A V4"
    style_header(ws)

    wrap = Alignment(vertical="top", wrap_text=True)
    thin = Border(
        left=Side(style="thin", color="E8EAED"),
        right=Side(style="thin", color="E8EAED"),
        top=Side(style="thin", color="E8EAED"),
        bottom=Side(style="thin", color="E8EAED"),
    )

    for r_idx, row in enumerate(rows, start=2):
        values = [
            row.get("id", ""),
            row.get("topic", ""),
            row.get("keywords", ""),
            row.get("sampleQuestion", ""),
            row.get("userId", ""),
            row.get("profile", ""),
            row.get("assignmentType", ""),
            row.get("hostCountry", ""),
            row.get("fullResponse", ""),
            row.get("followUp1", ""),
            row.get("followUp2", ""),
            row.get("followUp3", ""),
            row.get("validationStatus", "To review"),
            row.get("notes", ""),
        ]
        for c_idx, val in enumerate(values, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.alignment = wrap
            cell.border = thin
        ws.row_dimensions[r_idx].height = 120

    for i, w in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"

    # Summary sheet
    meta = wb.create_sheet("README")
    meta["A1"] = "Expat Concierge V4 — Offline Q&A Catalog"
    meta["A1"].font = Font(bold=True, size=14, color="0F3D4C")
    lines = [
        "",
        f"Generated: {data.get('generatedAt', '—')}",
        f"Rows: {len(rows)} topic handlers (+ variants)",
        "",
        "Purpose: Validate and refine offline chat answers before tester rollout.",
        "Source code: src/services/v4MobilityPrompt.js",
        "",
        "How to regenerate:",
        "  1. node scripts/export_offline_qa_catalog.mjs",
        "  2. python3 scripts/export_offline_qa_xlsx.py",
        "",
        "Trigger keywords syntax (column C):",
        "  • Use / or , for OR within a group: housing/rent/apartment",
        "  • Use + for AND between groups: need more/what if + housing/rent",
        "  • Set Validation status = Disabled to hide a row",
        "",
        "The live app reads this Excel file on each chat (auto-reload on save).",
    ]
    for i, line in enumerate(lines, 2):
        meta.cell(row=i, column=1, value=line)
    meta.column_dimensions["A"].width = 72

    XLSX_OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(XLSX_OUT)
    print(f"Wrote {len(rows)} rows → {XLSX_OUT}")


if __name__ == "__main__":
    main()
