#!/usr/bin/env python3
"""
Ingest Assignment Risk Matrix from Excel → JSON for V4 (and V3).

Reads business review overrides (yellow columns) when present.
Default inputs (first found wins unless --input is set):
  - OneDrive: Expat/expat-concierge/data/V3_Assignment_Risk_Matrix_v*.xlsx
  - docs/V*_Assignment_Risk_Matrix_*.xlsx

Outputs:
  - data/risk-matrix.json
  - assets/risk-matrix.json  (served at /assets/risk-matrix.json)
"""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
ONEDRIVE_DATA = Path(
    "/Users/renatobadaro/Library/CloudStorage/OneDrive-Anheuser-BuschInBev"
    "/Expat/expat-concierge/data"
)
DATA_OUT = ROOT / "data" / "risk-matrix.json"
ASSETS_OUT = ROOT / "assets" / "risk-matrix.json"

VALID_TIERS = {"blocker", "critical", "high", "medium", "low"}
VALID_GATES = {"start", "move", "settle", "complete", "support"}
MAX_SCORE_INFINITY = 99999


def find_default_xlsx() -> Path | None:
    candidates: list[Path] = []
    for base in [ONEDRIVE_DATA, ROOT / "docs"]:
        if not base.exists():
            continue
        candidates.extend(sorted(base.glob("*Assignment_Risk_Matrix*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True))
    return candidates[0] if candidates else None


def norm_tier(val: str | None) -> str | None:
    if not val:
        return None
    s = str(val).strip().lower()
    return s if s in VALID_TIERS else None


def norm_gate(val: str | None) -> str | None:
    if not val:
        return None
    s = str(val).strip().lower()
    return s if s in VALID_GATES else None


def parse_max_score(val) -> int:
    if val is None:
        return MAX_SCORE_INFINITY
    if isinstance(val, str) and val.strip() in {"∞", "inf", "infinity"}:
        return MAX_SCORE_INFINITY
    try:
        n = int(float(val))
        return n if n > 0 else MAX_SCORE_INFINITY
    except (TypeError, ValueError):
        return MAX_SCORE_INFINITY


def ingest(xlsx_path: Path) -> dict:
    wb = load_workbook(xlsx_path, data_only=True)
    ov = wb["Overview"]

    matrix: dict = {
        "version": str(ov["B3"].value or "1.0").strip(),
        "updated": str(ov["B4"].value or "").strip(),
        "purpose": str(ov["B5"].value or "").strip(),
        "sourceFile": xlsx_path.name,
        "tiers": {},
        "gates": {},
        "statusMultipliers": {},
        "levelThresholds": [],
        "escalationRules": [],
        "tasks": {},
    }

    ws_gates = wb["Gate Types"]
    for r in range(2, ws_gates.max_row + 1):
        key = ws_gates.cell(r, 1).value
        desc = ws_gates.cell(r, 2).value
        if key:
            matrix["gates"][str(key).strip()] = str(desc or "").strip()

    ws_tiers = wb["Tier Weights"]
    for r in range(2, ws_tiers.max_row + 1):
        key = ws_tiers.cell(r, 1).value
        if not key:
            continue
        key = str(key).strip()
        weight = ws_tiers.cell(r, 5).value or ws_tiers.cell(r, 3).value
        matrix["tiers"][key] = {
            "label": str(ws_tiers.cell(r, 2).value or key).strip(),
            "weight": int(float(weight or 2)),
            "desc": str(ws_tiers.cell(r, 4).value or "").strip(),
        }

    ws_levels = wb["Risk Levels"]
    for r in range(2, ws_levels.max_row + 1):
        level = ws_levels.cell(r, 1).value
        if not level:
            continue
        max_raw = ws_levels.cell(r, 4).value or ws_levels.cell(r, 3).value
        matrix["levelThresholds"].append({
            "level": str(level).strip(),
            "label": str(ws_levels.cell(r, 2).value or level).strip(),
            "maxScore": parse_max_score(max_raw),
        })

    ws_mult = wb["Status Multipliers"]
    for r in range(2, ws_mult.max_row + 1):
        status = ws_mult.cell(r, 1).value
        if not status:
            continue
        mult = ws_mult.cell(r, 3).value or ws_mult.cell(r, 2).value
        matrix["statusMultipliers"][str(status).strip()] = float(mult or 1.0)

    ws_rules = wb["Escalation Rules"]
    for r in range(3, ws_rules.max_row + 1):
        rule = ws_rules.cell(r, 1).value
        if rule and str(rule).strip():
            matrix["escalationRules"].append(str(rule).strip())

    ws_tasks = wb["Task Matrix"]
    changes_from_business = 0
    for r in range(2, ws_tasks.max_row + 1):
        tid = ws_tasks.cell(r, 2).value
        if tid is None:
            continue
        tid = int(tid)
        tier = norm_tier(ws_tasks.cell(r, 13).value) or norm_tier(ws_tasks.cell(r, 8).value) or "medium"
        gate = norm_gate(ws_tasks.cell(r, 14).value) or norm_gate(ws_tasks.cell(r, 10).value) or "support"
        note = ws_tasks.cell(r, 15).value or ws_tasks.cell(r, 12).value or ""
        if ws_tasks.cell(r, 13).value or ws_tasks.cell(r, 14).value or ws_tasks.cell(r, 15).value:
            changes_from_business += 1
        matrix["tasks"][tid] = {
            "tier": tier,
            "gate": gate,
            "note": str(note).strip(),
        }

    matrix["businessOverridesApplied"] = changes_from_business
    return matrix


def write_outputs(matrix: dict) -> None:
    DATA_OUT.parent.mkdir(parents=True, exist_ok=True)
    ASSETS_OUT.parent.mkdir(parents=True, exist_ok=True)
    payload = json.dumps(matrix, indent=2, ensure_ascii=False) + "\n"
    DATA_OUT.write_text(payload, encoding="utf-8")
    ASSETS_OUT.write_text(payload, encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description="Ingest risk matrix Excel to JSON")
    parser.add_argument("--input", type=Path, help="Path to Assignment Risk Matrix .xlsx")
    args = parser.parse_args()

    xlsx = args.input or find_default_xlsx()
    if not xlsx or not xlsx.exists():
        raise SystemExit("No risk matrix xlsx found. Pass --input path/to/file.xlsx")

    matrix = ingest(xlsx)
    write_outputs(matrix)

    print(f"Ingested: {xlsx}")
    print(f"  version: {matrix['version']}  updated: {matrix['updated']}")
    print(f"  tasks: {len(matrix['tasks'])}  business overrides: {matrix['businessOverridesApplied']}")
    print(f"  → {DATA_OUT}")
    print(f"  → {ASSETS_OUT}")


if __name__ == "__main__":
    main()
